"""Импорт расписания из CSV/XLSX. Формат описан в docs/IMPORT_FORMAT.md.

Колонки (заголовки на русском или английском, порядок любой):
группа | дата | пара | начало | конец | предмет | преподаватель | аудитория | тип | подгруппа
Обязательные: группа, дата, начало, конец, предмет. Если заголовка нет, берётся порядок выше.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time

from app.models import LessonRec

FIELD_ALIASES = {
    "group": ("группа", "group", "гр"),
    "day": ("дата", "date", "день"),
    "num": ("пара", "№", "номер", "num", "n"),
    "start": ("начало", "start", "с", "время начала"),
    "end": ("конец", "end", "по", "окончание", "время окончания"),
    "subject": ("предмет", "дисциплина", "subject"),
    "teacher": ("преподаватель", "препод", "teacher", "фио преподавателя"),
    "room": ("аудитория", "ауд", "кабинет", "room"),
    "kind": ("тип", "вид", "вид занятия", "kind"),
    "subgroup": ("подгруппа", "subgroup"),
}
DEFAULT_ORDER = ["group", "day", "num", "start", "end", "subject", "teacher", "room", "kind", "subgroup"]
REQUIRED = ("group", "day", "start", "end", "subject")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_day(s: str) -> str | None:
    s = s.strip()
    for pat in (r"^(\d{4})-(\d{1,2})-(\d{1,2})$",):
        m = re.match(pat, s)
        if m:
            y, mo, d = map(int, m.groups())
            try:
                return date(y, mo, d).isoformat()
            except ValueError:
                return None
    m = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$", s)
    if m:
        d, mo, y = map(int, m.groups())
        y += 2000 if y < 100 else 0
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return None
    return None


def parse_time(s: str) -> str | None:
    m = re.match(r"^(\d{1,2})[:.](\d{2})(?::\d{2})?$", s.strip())
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if not (0 <= h < 24 and 0 <= mi < 60):
        return None
    return f"{h:02d}:{mi:02d}"


def _map_header(row: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for i, cell in enumerate(row):
        c = cell.strip().lower()
        for field, names in FIELD_ALIASES.items():
            if c in names and field not in mapping:
                mapping[field] = i
    return mapping if all(f in mapping for f in REQUIRED) else None


def parse_rows(rows: list[list[str]]) -> tuple[list[LessonRec], list[str]]:
    """Возвращает (записи, ошибки). При наличии ошибок вызывающий код ничего не применяет."""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], ["Файл пустой."]
    mapping = _map_header(rows[0])
    start_idx = 1
    if mapping is None:
        mapping = {f: i for i, f in enumerate(DEFAULT_ORDER)}
        start_idx = 0
        first_day = parse_day(rows[0][1]) if len(rows[0]) > 1 else None
        if first_day is None:
            return [], ["Не найден заголовок. Нужны колонки: группа, дата, начало, конец, предмет."]
    out: list[LessonRec] = []
    errors: list[str] = []
    seen_num: dict[tuple[str, str, str], int] = {}
    for n, row in enumerate(rows[start_idx:], start=start_idx + 1):
        def get(field: str) -> str:
            i = mapping.get(field)
            return row[i].strip() if i is not None and i < len(row) else ""

        group, day_s, subject = get("group"), get("day"), get("subject")
        day = parse_day(day_s)
        start, end = parse_time(get("start")), parse_time(get("end"))
        problems = []
        if not group:
            problems.append("нет группы")
        if day is None:
            problems.append(f"неверная дата «{day_s}»")
        if start is None or end is None:
            problems.append("неверное время (нужно ЧЧ:ММ)")
        elif (end <= start):
            problems.append("конец раньше начала")
        if not subject:
            problems.append("нет предмета")
        num_s = get("num")
        if num_s:
            try:
                num = int(num_s)
            except ValueError:
                num = -1
                problems.append(f"неверный номер пары «{num_s}»")
        else:
            key = (group.upper(), day or "", "")
            seen_num[key] = seen_num.get(key, 0) + 1
            num = seen_num[key]
        sub_s = get("subgroup")
        sub: int | None = None
        if sub_s:
            if sub_s.isdigit() and sub_s in ("1", "2", "3"):
                sub = int(sub_s)
            else:
                problems.append(f"неверная подгруппа «{sub_s}»")
        if problems:
            errors.append(f"Строка {n}: " + "; ".join(problems))
            if len(errors) >= 10:
                errors.append("…остальные ошибки скрыты.")
                break
            continue
        out.append(LessonRec(group=group, day=day, num=num, start=start, end=end, subject=subject,  # type: ignore[arg-type]
                             teacher=get("teacher"), room=get("room"), kind=get("kind"), subgroup=sub))
    return out, errors


def read_csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не удалось определить кодировку файла (нужна UTF-8 или Windows-1251).")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"  # type: ignore[misc]
    return [[c for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [[_s(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def parse_file(filename: str, data: bytes) -> tuple[list[LessonRec], list[str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = read_xlsx(data)
    elif name.endswith((".csv", ".txt")):
        rows = read_csv(data)
    else:
        raise ValueError("Поддерживаются файлы .xlsx и .csv")
    return parse_rows(rows)


def parse_manual_line(line: str, group: str, today: date) -> tuple[LessonRec | None, str]:
    """Одна пара вручную: «дата;начало;конец;предмет;преподаватель;аудитория;тип» (последние три можно опустить)."""
    parts = [p.strip() for p in re.split(r"[;|]", line)]
    if len(parts) < 4:
        return None, "Нужно минимум: дата;начало;конец;предмет"
    day = parse_day(parts[0]) or None
    if day is None:
        m = re.match(r"^(\d{1,2})[./](\d{1,2})$", parts[0])
        if m:
            try:
                day = date(today.year, int(m.group(2)), int(m.group(1))).isoformat()
            except ValueError:
                day = None
    if day is None:
        return None, f"Неверная дата «{parts[0]}»"
    start, end = parse_time(parts[1]), parse_time(parts[2])
    if not start or not end or end <= start:
        return None, "Неверное время (ЧЧ:ММ, конец позже начала)"
    if not parts[3]:
        return None, "Не указан предмет"
    extra = parts[4:] + [""] * 3
    return LessonRec(group, day, 0, start, end, parts[3], extra[0], extra[1], extra[2]), ""
